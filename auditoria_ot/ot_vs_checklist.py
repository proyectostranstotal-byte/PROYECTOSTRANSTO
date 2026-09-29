"""Filtra las OT de la flota que tocan algún punto del checklist de auditoría.

Uso: python ot_vs_checklist.py <carpeta_exports> <desde> <hasta> [salida.xlsx]
Cada tarea y cada repuesto de la OT se clasifica por separado contra el checklist; la OT entra
si al menos un renglón corresponde a un punto auditable. Se muestran solo esos renglones.
"""
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from analisis_maestros import FLOTA, norm, vehiculos
from reparaciones_por_ot import ESTADOS, MESES, cantidad, leer_repuestos, leer_tareas

# (sección del checklist, punto, regex sobre el renglón normalizado, regex de exclusión)
CHECKLIST = [
    ("Documentación unidades", "Revisión técnica obligatoria – R.T.O.", r"\bVTV\b|\bRTO\b|REVISION TECNICA", None),
    ("Documentación unidades", "Chapa patente legible", r"PATENTE", None),
    ("Elementos cisterna", "Rótulo riesgo producto (rombos) / cartelería", r"ROMBO|CARTELER", None),
    ("Elementos cisterna", "N° ONU producto (panel naranja)", r"\bONU\b|PANEL NARANJA", None),
    ("Elementos cisterna", "Porta mangueras con 2 mangueras", r"PORTA ?MANGUERA", None),
    ("Elementos cisterna", "Estado mangueras, conexiones, bridas",
     r"MANGUERA.*(?:DESCARGA|PRODUCTO|CISTERN|CRIOG)|\bBRIDA", r"VIGIA|FUELLE|VEJIGA"),
    ("Elementos cisterna", "Sistema Stop-away", r"STOP ?-?AWAY", None),
    ("Elementos cisterna", "Gabinete libre de aceite, trapos", r"GABINETE", None),
    ("Elementos cisterna", "Funcionamiento de frenos",
     r"FRENO|PASTILLA|ZAPATA|PULMON|CAMPANA|\bABS\b", r"AUTOFRENANTE"),
    ("Elementos cisterna", "Alarma sonora y luces de retroceso", r"RETROCESO|ALARMA|SIRENA", None),
    ("Elementos cisterna", "Auxiliares", r"\bAUX", r"FARO|ESPEJO|LAMP|OPTICA"),
    ("Elementos cisterna", "Candado portamanguera", r"CANDADO", None),
    ("Estado y elementos unidad", "Estado general y limpieza", r"LAVAD|LIMPIEZA (?:GENERAL|DE (?:LA )?UNIDAD)", None),
    ("Estado y elementos unidad", "Matafuegos", r"MATA ?FUEGO|EXTINT", None),
    ("Estado y elementos unidad", "Cinturón de seguridad", r"CINTURON", None),
    ("Estado y elementos unidad", "Botiquín", r"BOTIQUIN", None),
    ("Estado y elementos unidad", "Arrestallamas", r"ARRESTA ?LLAMA", None),
    ("Estado y elementos unidad", "Calcomanía banda reflectiva", r"REFLECTIV|CALCOMAN|BANDA REFL", None),
    ("Estado y elementos unidad", "Calzas (2) y conos (6)", r"\bCALZAS?\b|CONOS? (?:DE )?SENAL", None),
    ("Estado y elementos unidad", "Funcionamiento luces",
     r"LAMP|FARO|OPTICA|\bLUZ\b|\bLUCES\b|\bMICA|\bICA\b|ICA ?FALTANTE|BALIZA|DESTELLADOR|STROBOS|REFLECTOR|TRILLER|\bAS DE LUZ",
     r"REFLECTIV"),
    ("Estado y elementos unidad", "Números de teléfono ante emergencias", r"TELEFONO", None),
    ("Estado y elementos unidad", "Control funcionamiento Michelin", r"MICHELIN", None),
    ("Estado y elementos unidad", "Botón de Imseg", r"IMSEG|PANICO", None),
    ("Estado y elementos unidad", "Espejos", r"ESPEJO", None),
    ("Documentación conductor", "Tacógrafo", r"TACOGRAF", None),
]


def puntos(linea):
    t = norm(linea)
    return [(sec, p) for sec, p, rx, excl in CHECKLIST
            if re.search(rx, t) and not (excl and re.search(excl, t))]


def main(carpeta, desde, hasta, salida):
    carpeta = Path(carpeta)
    o = pd.read_excel(carpeta / "ordenes_completo.xls", header=5, engine="openpyxl").dropna(subset=["Id OT"])
    o["pat"] = o["Patente"].astype(str).str.replace(" ", "").str.upper()
    o = o[o["pat"].isin(FLOTA) & (o["F.Ingreso"] >= pd.Timestamp(desde))
          & (o["F.Ingreso"] <= pd.Timestamp(hasta))].sort_values(["F.Ingreso", "Id OT"])
    v = vehiculos(carpeta / "vehiculos.xls").set_index("pat")
    tareas = leer_tareas(carpeta / "tareas_ot.txt")
    repuestos = leer_repuestos(carpeta / "items_remitos.txt")

    def unidad(p):
        modelo, interno = str(v.loc[p, "modelo"]), v.loc[p, "interno_modelo"]
        if modelo.startswith("CISTERNA"):
            return f"Cisterna int. {interno}"
        return modelo.replace("MB ", "") + (f" (int. {interno})" if isinstance(interno, str) else "")

    ots, renglones = [], []
    for _, r in o.iterrows():
        ot = int(r["Id OT"])
        cb = str(r["Cbte.Venta"]) if pd.notna(r["Cbte.Venta"]) else ""
        rep = repuestos.get(cb, {}) if cb != "RS-0000-00000000" else {}
        lineas = [("Tarea", f"{s}: {d}", d) for s, d, _ in tareas.get(ot, [])]
        lineas += [("Repuesto", f"{cantidad(c)} × {d}", d) for d, c in sorted(rep.items()) if c > 0]
        base = {"ot": ot, "fecha": r["F.Ingreso"].date(), "pat": r["pat"], "unidad": unidad(r["pat"]),
                "estado": ESTADOS.get(r["Estado"], r["Estado"])}
        rel_t, rel_r, pts = [], [], []
        for tipo, texto, desc in lineas:
            pp = puntos(desc)
            if not pp:
                continue
            (rel_t if tipo == "Tarea" else rel_r).append(texto)
            pts += [p for _, p in pp]
            for sec, p in pp:
                renglones.append({**base, "seccion": sec, "punto": p, "tipo": tipo, "detalle": texto})
        if pts:
            ots.append({**base, "puntos": ", ".join(dict.fromkeys(pts)),
                        "tareas": "\n".join(f"• {x}" for x in rel_t) or "—",
                        "repuestos": "\n".join(f"• {x}" for x in rel_r) or "—"})
    d, rg = pd.DataFrame(ots), pd.DataFrame(renglones)
    escribir(d, rg, len(o), salida)
    return d, rg, len(o)


def escribir(d, rg, total, salida):
    H, HF = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
    GRUPO = PatternFill("solid", fgColor="DDEBF7")
    ROJO = PatternFill("solid", fgColor="FCE4D6")
    linea = Border(bottom=Side(style="thin", color="BFBFBF"))
    arriba = Alignment(vertical="top", wrap_text=True)

    def encabezado(ws, titulo, sub, cols, anchos):
        ws["A1"], ws["A2"] = titulo, sub
        ws["A1"].font, ws["A2"].font = Font(bold=True, size=14), Font(italic=True, color="595959")
        ws.append([])
        ws.append(cols)
        for c in ws[4]:
            c.font, c.fill = H, HF
        for i, w in enumerate(anchos):
            ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes = "A5"

    def fila(ws, valores, col_fecha=None, col_estado=None, estado=None):
        ws.append(valores)
        n = ws.max_row
        for c in ws[n]:
            c.alignment, c.border = arriba, linea
        if col_fecha:
            ws.cell(n, col_fecha).number_format = "DD/MM/YYYY"
        if col_estado and estado == "En diagnóstico":
            ws.cell(n, col_estado).fill = ROJO

    def grupo(ws, texto, ncols):
        ws.append([texto])
        for i in range(1, ncols + 1):
            ws.cell(ws.max_row, i).fill = GRUPO
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)

    wb = Workbook()
    wb.remove(wb.active)
    sub = (f"{len(d)} de {total} OT de la flota tocan algún punto del checklist. "
           "Se muestran solo las tareas y repuestos relacionados.")

    # 1) Por punto del checklist
    ws = wb.create_sheet("Por punto de auditoría")
    encabezado(ws, "OT relacionadas con el checklist de auditoría — por punto", sub,
               ["Punto del checklist", "OT N°", "Fecha", "Patente", "Unidad", "Estado", "Qué se hizo / qué se puso"],
               [34, 8, 11, 10, 22, 13, 70])
    orden = [(s, p) for s, p, _, _ in CHECKLIST]
    for sec, p in orden:
        x = rg[rg.punto == p]
        if x.empty:
            continue
        grupo(ws, f"{sec} › {p}  ({x.ot.nunique()} OT)", 7)
        for (ot, fecha, pat, uni, est), y in x.groupby(["ot", "fecha", "pat", "unidad", "estado"], sort=False):
            fila(ws, ["", ot, fecha, pat, uni, est, "\n".join(f"• {t}" for t in dict.fromkeys(y.detalle))],
                 col_fecha=3, col_estado=6, estado=est)
    sin = [p for s, p in orden if p not in set(rg.punto)]
    ws.append([])
    ws.append([f"Sin OT en el período: {', '.join(sin)}."])
    ws.cell(ws.max_row, 1).font = Font(italic=True, color="595959")

    # 2) Por unidad
    ws = wb.create_sheet("Por unidad")
    encabezado(ws, "OT relacionadas con el checklist — por unidad", sub,
               ["Patente / Unidad", "OT N°", "Fecha", "Estado", "Puntos del checklist",
                "Tareas relacionadas", "Repuestos relacionados"],
               [34, 8, 11, 13, 30, 55, 50])
    for pat in d.groupby("pat").size().sort_values(ascending=False).index:
        x = d[d.pat == pat]
        grupo(ws, f"{pat} — {x.unidad.iloc[0]}  ({len(x)} OT)", 7)
        for _, r in x.iterrows():
            fila(ws, ["", r.ot, r.fecha, r.estado, r.puntos, r.tareas, r.repuestos],
                 col_fecha=3, col_estado=4, estado=r.estado)

    # 3) Una hoja por mes, lista plana con filtros
    for anio, mes in sorted({(f.year, f.month) for f in d.fecha}):
        x = d[[f.year == anio and f.month == mes for f in d.fecha]]
        ws = wb.create_sheet(MESES[mes])
        encabezado(ws, f"OT relacionadas con el checklist — {MESES[mes]} {anio}", f"{len(x)} OT",
                   ["OT N°", "Fecha", "Patente", "Unidad", "Estado", "Puntos del checklist",
                    "Tareas relacionadas", "Repuestos relacionados"],
                   [8, 11, 10, 22, 13, 30, 55, 50])
        for _, r in x.iterrows():
            fila(ws, [r.ot, r.fecha, r.pat, r.unidad, r.estado, r.puntos, r.tareas, r.repuestos],
                 col_fecha=2, col_estado=5, estado=r.estado)
        ws.auto_filter.ref = f"A4:H{ws.max_row}"

    # 4) Resumen
    ws = wb.create_sheet("Resumen", 0)
    encabezado(ws, "Resumen: OT por punto del checklist", sub,
               ["Sección", "Punto del checklist", "OT", "Unidades", "Patentes"], [26, 40, 6, 9, 80])
    for sec, p in orden:
        x = rg[rg.punto == p]
        fila(ws, [sec, p, x.ot.nunique(), x.pat.nunique(), ", ".join(sorted(x.pat.unique()))])
        if x.empty:
            for c in ws[ws.max_row]:
                c.font = Font(color="A6A6A6")
    wb.save(salida)


if __name__ == "__main__":
    a = sys.argv
    d, rg, total = main(a[1], a[2], a[3], a[4] if len(a) > 4 else "ot_checklist.xlsx")
    print(f"{len(d)} de {total} OT relacionadas")
    print(rg.groupby("punto").ot.nunique().sort_values(ascending=False).to_string())
