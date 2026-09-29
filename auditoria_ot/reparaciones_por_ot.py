"""Reparaciones por OT de la flota auditada: qué se le hizo a cada unidad y qué repuestos se le pusieron.

Uso: python reparaciones_por_ot.py <carpeta> <desde AAAA-MM-DD> <hasta AAAA-MM-DD> [salida.xlsx]

La carpeta debe tener los exports de Copérnico:
  ordenes_completo.xls  (exportado SIN "formato rápido"; trae Cbte.Venta)
  tareas_ot.txt         (mano de obra por OT)
  items_remitos.txt     (movimientos de stock; las salidas RSM son los repuestos de cada OT)
  vehiculos.xls         (para el interno de las cisternas)
Vínculo OT <-> repuestos: Cbte.Venta RS-2025-00000821  <->  Nro. Cbte. RSM-R-2025-00000821.
"""
import re
import sys
from collections import defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from analisis_maestros import FLOTA, ITEMS, norm, vehiculos

ESTADOS = {"Cerrada ( No Permitira Modif)": "Cerrada", "Diagnostico": "En diagnóstico",
           "Terminado (Antes Facturar)": "Terminada"}
MESES = ["", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto",
         "Septiembre", "Octubre", "Noviembre", "Diciembre"]


def entero(s):
    return int(str(s).replace(",", "").replace(".", "").strip() or 0)


def numero(s):
    s = str(s).strip()
    return float(s.replace(",", "")) if s else 0.0


def leer_txt(path):
    return [linea.split(";") for linea in
            Path(path).read_text(encoding="cp1252").splitlines()[1:] if linea.strip()]


def leer_tareas(path):
    tareas = defaultdict(list)
    for f in leer_txt(path):
        if len(f) > 27:  # un ";" dentro de la descripción: se vuelve a unir
            extra = len(f) - 27
            f = f[:15] + [";".join(f[15:16 + extra])] + f[16 + extra:]
        sector = f[13].strip().title()
        desc = f[15].strip()
        if desc:
            tareas[entero(f[0])].append((sector, desc, f[16].strip()))
    return tareas


def leer_repuestos(path):
    rep = defaultdict(lambda: defaultdict(float))
    for f in leer_txt(path):
        if f[3] != "RSM":
            continue
        m = re.match(r"RSM-R-(\d{4})-(\d+)", f[4].strip())
        if m:
            rep[f"RS-{m[1]}-{m[2]}"][f[6].strip()] += -numero(f[8])
    return rep


def cantidad(x):
    return f"{x:g}".replace(".", ",")


def items_auditoria(texto):
    t = norm(texto)
    return [item for item, rx in ITEMS.items() if re.search(rx, t)]


def main(carpeta, desde, hasta, salida):
    carpeta = Path(carpeta)
    o = pd.read_excel(carpeta / "ordenes_completo.xls", header=5, engine="openpyxl")
    o = o.dropna(subset=["Id OT"])
    o["pat"] = o["Patente"].astype(str).str.replace(" ", "").str.upper()
    o = o[o["pat"].isin(FLOTA)
          & (o["F.Ingreso"] >= pd.Timestamp(desde)) & (o["F.Ingreso"] <= pd.Timestamp(hasta))]
    o = o.sort_values(["F.Ingreso", "Id OT"])

    v = vehiculos(carpeta / "vehiculos.xls").set_index("pat")
    tareas = leer_tareas(carpeta / "tareas_ot.txt")
    repuestos = leer_repuestos(carpeta / "items_remitos.txt")

    def unidad(p):
        modelo, interno = str(v.loc[p, "modelo"]), v.loc[p, "interno_modelo"]
        if modelo.startswith("CISTERNA"):
            return f"Cisterna int. {interno}"
        return modelo.replace("MB ", "") + (f" (int. {interno})" if isinstance(interno, str) else "")

    filas = []
    for _, r in o.iterrows():
        ot = int(r["Id OT"])
        tt = tareas.get(ot, [])
        cb = str(r["Cbte.Venta"]) if pd.notna(r["Cbte.Venta"]) else ""
        rr = {d: c for d, c in repuestos.get(cb, {}).items() if c > 0} if cb != "RS-0000-00000000" else {}
        txt_t = "\n".join(f"• {s}: {d}" + ("" if e in ("Realizada", "") else f" ({e.lower()})")
                          for s, d, e in tt)
        txt_r = "\n".join(f"• {cantidad(c)} × {d}" for d, c in sorted(rr.items()))
        fin = r["Fecha Finalizacion"]
        filas.append({
            "ot": ot, "fecha": r["F.Ingreso"].date(), "pat": r["pat"], "unidad": unidad(r["pat"]),
            "estado": ESTADOS.get(r["Estado"], r["Estado"]),
            "fin": fin.date() if pd.notna(fin) and fin.year > 2000 else None,
            "trabajos": txt_t or "(sin tareas cargadas)", "repuestos": txt_r or "—",
            "auditoria": ", ".join(dict.fromkeys(items_auditoria(" ".join(d for _, d, _ in tt)
                                                                  + " " + " ".join(rr)))),
            "n_tareas": len(tt), "n_rep": len(rr),
        })
    d = pd.DataFrame(filas)
    escribir(d, salida)
    return d


def escribir(d, salida):
    H, HF = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
    ROJO = PatternFill("solid", fgColor="FCE4D6")
    linea = Border(bottom=Side(style="thin", color="BFBFBF"))
    arriba = Alignment(vertical="top", wrap_text=True)
    cols = ["OT N°", "Fecha", "Patente", "Unidad", "Estado", "Qué se le hizo (tareas)",
            "Repuestos que se le pusieron", "Ítems de auditoría relacionados"]
    anchos = [8, 11, 10, 22, 13, 60, 55, 30]

    wb = Workbook()
    wb.remove(wb.active)
    meses = sorted({(f.year, f.month) for f in d["fecha"]})
    for anio, mes in meses:
        sub = d[[f.year == anio and f.month == mes for f in d["fecha"]]]
        ws = wb.create_sheet(MESES[mes])
        ws["A1"] = f"Reparaciones de la flota auditada — {MESES[mes]} {anio}"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = (f"{len(sub)} OT · {sub['n_rep'].gt(0).sum()} con repuestos · "
                    f"{(sub['estado'] == 'En diagnóstico').sum()} todavía en diagnóstico (en rojo)")
        ws["A2"].font = Font(italic=True, color="595959")
        ws.append([])
        ws.append(cols)
        for c in ws[4]:
            c.font, c.fill = H, HF
        for _, r in sub.iterrows():
            ws.append([r.ot, r.fecha, r.pat, r.unidad, r.estado, r.trabajos, r.repuestos, r.auditoria])
            fila = ws.max_row
            ws.cell(fila, 2).number_format = "DD/MM/YYYY"
            for c in ws[fila]:
                c.alignment, c.border = arriba, linea
            if r.estado == "En diagnóstico":
                ws.cell(fila, 5).fill = ROJO
        for i, w in enumerate(anchos):
            ws.column_dimensions["ABCDEFGH"[i]].width = w
        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:H{ws.max_row}"

    # Ficha por unidad: todas sus OT del período una debajo de la otra
    ws = wb.create_sheet("Por unidad", 0)
    ws["A1"] = "Qué se le hizo a cada unidad"
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Patente / Unidad", "OT N°", "Fecha", "Estado", "Qué se le hizo (tareas)",
               "Repuestos que se le pusieron"])
    for c in ws[3]:
        c.font, c.fill = H, HF
    orden = d.groupby("pat").size().sort_values(ascending=False).index
    for p in orden:
        x = d[d.pat == p]
        ws.append([f"{p} — {x.unidad.iloc[0]}  ({len(x)} OT)"])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor="DDEBF7")
        for i in range(2, 7):
            ws.cell(ws.max_row, i).fill = PatternFill("solid", fgColor="DDEBF7")
        for _, r in x.iterrows():
            ws.append(["", r.ot, r.fecha, r.estado, r.trabajos, r.repuestos])
            fila = ws.max_row
            ws.cell(fila, 3).number_format = "DD/MM/YYYY"
            for c in ws[fila]:
                c.alignment, c.border = arriba, linea
            if r.estado == "En diagnóstico":
                ws.cell(fila, 4).fill = ROJO
    for col, w in zip("ABCDEF", [34, 8, 11, 13, 60, 55]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    wb.save(salida)


if __name__ == "__main__":
    a = sys.argv
    d = main(a[1], a[2], a[3], a[4] if len(a) > 4 else "reparaciones_flota.xlsx")
    print(len(d), "OT;", d.n_tareas.eq(0).sum(), "sin tareas;", d.n_rep.eq(0).sum(), "sin repuestos")
